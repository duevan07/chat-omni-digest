from __future__ import annotations

import mimetypes
from pathlib import Path

from .models import Attachment, Conversation


def summarize_file(path: str | Path, max_chars: int = 4000) -> str:
    file_path = Path(path)
    mime_type, _ = mimetypes.guess_type(str(file_path))
    suffix = file_path.suffix.lower()

    if suffix in {".txt", ".md", ".csv", ".log", ".json"}:
        text = file_path.read_text(encoding="utf-8", errors="replace")
        return text[:max_chars]

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception:
            return "PDF file detected. Install chat-omni-digest[pdf] to extract text."
        reader = PdfReader(str(file_path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return text[:max_chars] or "PDF file detected, but no extractable text was found."

    if suffix == ".docx":
        try:
            from docx import Document  # type: ignore
        except Exception:
            return "Word document detected. Install chat-omni-digest[office] to extract text."
        doc = Document(str(file_path))
        text = "\n".join(p.text for p in doc.paragraphs)
        return text[:max_chars]

    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl  # type: ignore
        except Exception:
            return "Spreadsheet detected. Install chat-omni-digest[office] to extract text."
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets[:3]:
            lines.append(f"[{ws.title}]")
            for row in ws.iter_rows(max_row=20, values_only=True):
                lines.append("\t".join("" if v is None else str(v) for v in row))
        return "\n".join(lines)[:max_chars]

    return f"{mime_type or 'binary'} file: {file_path.name} ({file_path.stat().st_size} bytes)"


def enrich_attachments(conversation: Conversation, max_chars: int = 4000) -> Conversation:
    for message in conversation.messages:
        for attachment in message.attachments:
            if attachment.path and Path(attachment.path).exists() and attachment.kind == "file":
                attachment.summary = summarize_file(attachment.path, max_chars=max_chars)
    return conversation

